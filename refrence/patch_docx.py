from docx import Document
import re, os

OUTPUT_PATH = "/mnt/e/naveen/refrence/LastZone_DevOps_Report_DHARNIDHARAN_V.docx"

doc = Document(OUTPUT_PATH)

fix_map = {
    'canteen_data': 'lastzone_postgres_data',
    'Jayasree273/DigitalCanteen': 'urwishservice-max/lastone-tournament',
    'Jayasree273': 'urwishservice-max',
    'DigitalCanteen.git': 'lastone-tournament.git',
    'DigitalCanteen': 'lastone-tournament',
}

count = 0
for para in doc.paragraphs:
    for run in para.runs:
        orig = run.text
        for old, new in fix_map.items():
            run.text = run.text.replace(old, new)
        if run.text != orig:
            count += 1

print(f"Fixed {count} additional occurrences")
doc.save(OUTPUT_PATH)

# Final verification
doc2 = Document(OUTPUT_PATH)
all_text = ' '.join([p.text for p in doc2.paragraphs])
canteen_left = len(re.findall('canteen', all_text, re.IGNORECASE))
jayasree_left = len(re.findall('jayasree', all_text, re.IGNORECASE))
old_roll_left = len(re.findall('2303717620522022', all_text))
print("=== FINAL VERIFICATION ===")
print(f"Remaining 'canteen':  {canteen_left}")
print(f"Remaining 'JAYASREE': {jayasree_left}")
print(f"Remaining old roll:   {old_roll_left}")

if canteen_left > 0:
    for i, p in enumerate(doc2.paragraphs):
        if 'canteen' in p.text.lower():
            print(f"  Still at para {i}: {p.text[:100]}")

size = os.path.getsize(OUTPUT_PATH)
print(f"\nFile: {OUTPUT_PATH}")
print(f"Size: {size:,} bytes ({size/1024/1024:.1f} MB)")
print("Done!")
